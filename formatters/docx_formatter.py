from io import BytesIO
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from .utils import LOGO_PATH, sanitize_text, split_sections, terms_from_text


def _set_cell_shading(cell, fill="E8EEF7"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def format_docx(text: str, doc_type: str, source_terms: str = "") -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(11)

    if LOGO_PATH.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(LOGO_PATH), width=Inches(1.15))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(sanitize_text(doc_type).upper())
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    for heading, body in split_sections(text):
        p = doc.add_paragraph()
        r = p.add_run(heading)
        r.bold = True
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)
        for line in body:
            if not line:
                continue
            bp = doc.add_paragraph(line)
            bp.paragraph_format.space_after = Pt(5)

    terms = [item.strip() for item in source_terms.split(";") if item.strip()] if source_terms.strip() else terms_from_text(text)
    if terms:
        doc.add_paragraph()
        h = doc.add_paragraph()
        r = h.add_run("KEY TERMS")
        r.bold = True
        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "No."
        table.rows[0].cells[1].text = "Term / Clause"
        for c in table.rows[0].cells:
            _set_cell_shading(c)
            for rr in c.paragraphs[0].runs:
                rr.bold = True
        for i, term in enumerate(terms[:15], 1):
            cells = table.add_row().cells
            cells[0].text = str(i)
            cells[1].text = term

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.text = "LegalEase | AI-generated draft — professional legal review recommended"
    for run in footer.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(8)

    output = BytesIO()
    doc.save(output)
    return output.getvalue()
